"""Renders analytics_report_service rows to PDF (ReportLab) and Excel
(openpyxl). Kept separate from the query logic so the export format never
influences what data is queried (improvement spec #11/#12)."""
import io
from datetime import datetime, timezone

DEMAND_COLUMNS = [
    ("product", "Product"), ("category", "Category"), ("shop", "Shop"),
    ("views", "Views"), ("reserve_requests", "Reserve Requests"),
    ("completed_reservations", "Completed"), ("price", "Price"), ("availability", "Availability"),
]

SUPPLY_COLUMNS = [
    ("product", "Product"), ("category", "Category"), ("shop", "Shop"),
    ("price", "Price"), ("availability", "Availability"), ("rating", "Rating"),
    ("reviews", "Reviews"), ("last_updated", "Last Updated"),
]


def _format_cell(value):
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.2f}" if value % 1 else f"{value:.0f}"
    if isinstance(value, datetime):
        return value.strftime("%d %b %Y")
    return str(value)


def build_pdf(report_type, rows, start=None, end=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm

    columns = DEMAND_COLUMNS if report_type == "demand" else SUPPLY_COLUMNS
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4),
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()

    elements = [
        Paragraph("NearCart", styles["Title"]),
        Paragraph(f"{report_type.title()} Report", styles["Heading2"]),
        Paragraph(
            f"Generated {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M UTC')}"
            + (f" &nbsp;·&nbsp; Range: {start} to {end}" if start or end else ""),
            styles["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    header = [label for _, label in columns]
    data = [header]
    for row in rows:
        data.append([_format_cell(row.get(key)) for key, _ in columns])

    if len(data) == 1:
        elements.append(Paragraph("No matching data for this range.", styles["Normal"]))
    else:
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e7e5e4")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f4")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer


def build_excel(report_type, rows, start=None, end=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    columns = DEMAND_COLUMNS if report_type == "demand" else SUPPLY_COLUMNS
    wb = Workbook()
    ws = wb.active
    ws.title = f"{report_type.title()} Report"

    ws.append(["NearCart", f"{report_type.title()} Report"])
    ws.append([f"Generated {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}",
               f"Range: {start or 'all'} to {end or 'all'}"])
    ws.append([])

    header_row_index = 4
    ws.append([label for _, label in columns])
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for cell in ws[header_row_index]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for row in rows:
        ws.append([_format_cell(row.get(key)) for key, _ in columns])

    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
